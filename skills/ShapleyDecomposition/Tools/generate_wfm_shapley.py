#!/usr/bin/env python3
"""
Generate WFM Shapley Decomposition Template using openpyxl.

Creates a 12-month Excel workbook for staffing variance analysis using
Shapley decomposition. Three variance factors: Volume, AHT, Shrinkage.
Occupancy is held constant as a planning assumption.

Usage:
    python generate_wfm_shapley.py --output WFM_Shapley_Template.xlsx --year 2025
"""

import argparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Colors
COLORS = {
    'header': '0F172A',
    'header_text': 'FFFFFF',
    'input_blue': 'EFF6FF',
    'input_orange': 'FFF7ED',
    'section': '1D4ED8',
    'gold': 'F59E0B',
    'light_gray': 'F1F5F9',
}

MONTHS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']

def apply_header_style(cell):
    cell.font = Font(bold=True, color=COLORS['header_text'])
    cell.fill = PatternFill(start_color=COLORS['header'], end_color=COLORS['header'], fill_type='solid')
    cell.alignment = Alignment(horizontal='center', vertical='center')

def apply_section_style(cell):
    cell.font = Font(bold=True, size=12, color=COLORS['section'])

def apply_input_style(cell, input_type='plan'):
    color = COLORS['input_orange'] if input_type == 'actual' else COLORS['input_blue']
    cell.fill = PatternFill(start_color=color, end_color=color, fill_type='solid')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    cell.border = thin_border

def generate_wfm_shapley(output_path: str, year: int = 2024):
    wb = Workbook()
    ws = wb.active
    ws.title = "Shapley Analysis"

    # Column widths
    ws.column_dimensions['A'].width = 24
    for i in range(2, 14):
        ws.column_dimensions[get_column_letter(i)].width = 12

    row = 1

    # Title
    ws.merge_cells(f'A{row}:M{row}')
    ws[f'A{row}'] = f'WFM Staffing Variance Analysis - {year}'
    ws[f'A{row}'].font = Font(bold=True, size=16, color=COLORS['header'])
    ws[f'A{row}'].alignment = Alignment(horizontal='center')
    row += 2

    # Month headers
    ws[f'A{row}'] = 'Month'
    apply_header_style(ws[f'A{row}'])
    for m in range(12):
        cell = ws.cell(row=row, column=m+2)
        cell.value = f"{MONTHS[m]}-{str(year)[-2:]}"
        apply_header_style(cell)
    row += 2

    # ========== PLANNING ASSUMPTIONS ==========
    ws[f'A{row}'] = 'PLANNING ASSUMPTIONS'
    apply_section_style(ws[f'A{row}'])
    row += 1

    # Work Hours
    work_hours_row = row
    ws[f'A{row}'] = 'Work Hours/Month'
    for m in range(12):
        cell = ws.cell(row=row, column=m+2)
        cell.value = 160
        apply_input_style(cell, 'plan')
        cell.number_format = '#,##0'
    row += 1

    # Occupancy (CONSTANT)
    occ_row = row
    ws[f'A{row}'] = 'Planned Occupancy'
    ws[f'A{row}'].font = Font(bold=True)
    for m in range(12):
        cell = ws.cell(row=row, column=m+2)
        cell.value = 0.80
        apply_input_style(cell, 'plan')
        cell.number_format = '0.0%'
    row += 1

    ws[f'A{row}'] = '↑ Occupancy is a planning constant, not a variance driver'
    ws[f'A{row}'].font = Font(italic=True, size=10, color='64748B')
    row += 2

    # ========== VOLUME (CALLS) ==========
    ws[f'A{row}'] = 'VOLUME (CALLS)'
    apply_section_style(ws[f'A{row}'])
    row += 1

    plan_vol_row = row
    ws[f'A{row}'] = 'Plan/Budget'
    for m in range(12):
        cell = ws.cell(row=row, column=m+2)
        cell.value = 25000 + m * 100
        apply_input_style(cell, 'plan')
        cell.number_format = '#,##0'
    row += 1

    fcst_vol_row = row
    ws[f'A{row}'] = 'Forecast'
    for m in range(12):
        cell = ws.cell(row=row, column=m+2)
        cell.value = 26000 + m * 200
        apply_input_style(cell, 'forecast')
        cell.number_format = '#,##0'
    row += 1

    act_vol_row = row
    ws[f'A{row}'] = 'Actual'
    for m in range(12):
        cell = ws.cell(row=row, column=m+2)
        cell.value = 27000 + m * 150
        apply_input_style(cell, 'actual')
        cell.number_format = '#,##0'
    row += 1

    # Variance %
    ws[f'A{row}'] = 'Fcst vs Act %'
    for m in range(12):
        col = get_column_letter(m+2)
        cell = ws.cell(row=row, column=m+2)
        cell.value = f'=IFERROR(({col}{act_vol_row}-{col}{fcst_vol_row})/{col}{fcst_vol_row},0)'
        cell.number_format = '0.0%'
    row += 1

    ws[f'A{row}'] = 'Plan vs Act %'
    for m in range(12):
        col = get_column_letter(m+2)
        cell = ws.cell(row=row, column=m+2)
        cell.value = f'=IFERROR(({col}{act_vol_row}-{col}{plan_vol_row})/{col}{plan_vol_row},0)'
        cell.number_format = '0.0%'
    row += 2

    # ========== AHT ==========
    ws[f'A{row}'] = 'AHT (SECONDS)'
    apply_section_style(ws[f'A{row}'])
    row += 1

    plan_aht_row = row
    ws[f'A{row}'] = 'Plan/Budget'
    for m in range(12):
        cell = ws.cell(row=row, column=m+2)
        cell.value = 600
        apply_input_style(cell, 'plan')
        cell.number_format = '#,##0'
    row += 1

    fcst_aht_row = row
    ws[f'A{row}'] = 'Forecast'
    for m in range(12):
        cell = ws.cell(row=row, column=m+2)
        cell.value = 620 + m * 5
        apply_input_style(cell, 'forecast')
        cell.number_format = '#,##0'
    row += 1

    act_aht_row = row
    ws[f'A{row}'] = 'Actual'
    for m in range(12):
        cell = ws.cell(row=row, column=m+2)
        cell.value = 640 + m * 8
        apply_input_style(cell, 'actual')
        cell.number_format = '#,##0'
    row += 1

    ws[f'A{row}'] = 'Fcst vs Act %'
    for m in range(12):
        col = get_column_letter(m+2)
        cell = ws.cell(row=row, column=m+2)
        cell.value = f'=IFERROR(({col}{act_aht_row}-{col}{fcst_aht_row})/{col}{fcst_aht_row},0)'
        cell.number_format = '0.0%'
    row += 1

    ws[f'A{row}'] = 'Plan vs Act %'
    for m in range(12):
        col = get_column_letter(m+2)
        cell = ws.cell(row=row, column=m+2)
        cell.value = f'=IFERROR(({col}{act_aht_row}-{col}{plan_aht_row})/{col}{plan_aht_row},0)'
        cell.number_format = '0.0%'
    row += 2

    # ========== SHRINKAGE ==========
    ws[f'A{row}'] = 'SHRINKAGE'
    apply_section_style(ws[f'A{row}'])
    row += 1

    plan_shrink_row = row
    ws[f'A{row}'] = 'Plan/Budget'
    for m in range(12):
        cell = ws.cell(row=row, column=m+2)
        cell.value = 0.35
        apply_input_style(cell, 'plan')
        cell.number_format = '0.0%'
    row += 1

    fcst_shrink_row = row
    ws[f'A{row}'] = 'Forecast'
    for m in range(12):
        cell = ws.cell(row=row, column=m+2)
        cell.value = 0.38 + m * 0.005
        apply_input_style(cell, 'forecast')
        cell.number_format = '0.0%'
    row += 1

    act_shrink_row = row
    ws[f'A{row}'] = 'Actual'
    for m in range(12):
        cell = ws.cell(row=row, column=m+2)
        cell.value = 0.40 + m * 0.008
        apply_input_style(cell, 'actual')
        cell.number_format = '0.0%'
    row += 1

    ws[f'A{row}'] = 'Fcst vs Act %'
    for m in range(12):
        col = get_column_letter(m+2)
        cell = ws.cell(row=row, column=m+2)
        cell.value = f'=IFERROR(({col}{act_shrink_row}-{col}{fcst_shrink_row})/{col}{fcst_shrink_row},0)'
        cell.number_format = '0.0%'
    row += 1

    ws[f'A{row}'] = 'Plan vs Act %'
    for m in range(12):
        col = get_column_letter(m+2)
        cell = ws.cell(row=row, column=m+2)
        cell.value = f'=IFERROR(({col}{act_shrink_row}-{col}{plan_shrink_row})/{col}{plan_shrink_row},0)'
        cell.number_format = '0.0%'
    row += 2

    # ========== FTE REQUIREMENTS ==========
    ws[f'A{row}'] = 'FTE REQUIREMENTS'
    apply_section_style(ws[f'A{row}'])
    row += 1

    ws[f'A{row}'] = 'Formula: (Volume × AHT) / (WorkHours × 3600 × Occ × (1-Shrink))'
    ws[f'A{row}'].font = Font(italic=True, size=10)
    row += 1

    # Plan FTE
    plan_fte_row = row
    ws[f'A{row}'] = 'Plan Required FTE'
    ws[f'A{row}'].font = Font(bold=True)
    for m in range(12):
        col = get_column_letter(m+2)
        cell = ws.cell(row=row, column=m+2)
        cell.value = f'=IFERROR(({col}{plan_vol_row}*{col}{plan_aht_row})/({col}{work_hours_row}*3600*{col}{occ_row}*(1-{col}{plan_shrink_row})),0)'
        cell.number_format = '#,##0.00'
    row += 1

    # Forecast FTE
    fcst_fte_row = row
    ws[f'A{row}'] = 'Forecast Required FTE'
    ws[f'A{row}'].font = Font(bold=True)
    for m in range(12):
        col = get_column_letter(m+2)
        cell = ws.cell(row=row, column=m+2)
        cell.value = f'=IFERROR(({col}{fcst_vol_row}*{col}{fcst_aht_row})/({col}{work_hours_row}*3600*{col}{occ_row}*(1-{col}{fcst_shrink_row})),0)'
        cell.number_format = '#,##0.00'
    row += 1

    # Actual FTE
    act_fte_row = row
    ws[f'A{row}'] = 'Actual Required FTE'
    ws[f'A{row}'].font = Font(bold=True)
    for m in range(12):
        col = get_column_letter(m+2)
        cell = ws.cell(row=row, column=m+2)
        cell.value = f'=IFERROR(({col}{act_vol_row}*{col}{act_aht_row})/({col}{work_hours_row}*3600*{col}{occ_row}*(1-{col}{act_shrink_row})),0)'
        cell.number_format = '#,##0.00'
    row += 2

    # Variance Summary
    ws[f'A{row}'] = 'VARIANCE SUMMARY'
    apply_section_style(ws[f'A{row}'])
    row += 1

    fcst_var_row = row
    ws[f'A{row}'] = 'Fcst vs Actual (FTE)'
    ws[f'A{row}'].font = Font(bold=True)
    for m in range(12):
        col = get_column_letter(m+2)
        cell = ws.cell(row=row, column=m+2)
        cell.value = f'={col}{act_fte_row}-{col}{fcst_fte_row}'
        cell.number_format = '#,##0.00'
        cell.font = Font(bold=True)
    row += 1

    plan_var_row = row
    ws[f'A{row}'] = 'Plan vs Actual (FTE)'
    ws[f'A{row}'].font = Font(bold=True)
    for m in range(12):
        col = get_column_letter(m+2)
        cell = ws.cell(row=row, column=m+2)
        cell.value = f'={col}{act_fte_row}-{col}{plan_fte_row}'
        cell.number_format = '#,##0.00'
        cell.font = Font(bold=True)
    row += 3

    # ========== SHAPLEY: FORECAST vs ACTUAL ==========
    ws[f'A{row}'] = 'SHAPLEY: FORECAST vs ACTUAL'
    apply_section_style(ws[f'A{row}'])
    ws[f'A{row}'].fill = PatternFill(start_color=COLORS['gold'], end_color=COLORS['gold'], fill_type='solid')
    ws[f'A{row}'].font = Font(bold=True, size=12, color=COLORS['header'])
    row += 1

    ws[f'A{row}'] = 'Coalition values (baseline = forecast values)'
    ws[f'A{row}'].font = Font(italic=True, size=10)
    row += 1

    # Coalition values
    fcst_v0 = row
    ws[f'A{row}'] = 'v(∅) - All at Forecast'
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'={col}{fcst_fte_row}'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 1

    fcst_vV = row
    ws[f'A{row}'] = 'v(V) - Volume at Actual'
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'=IFERROR(({col}{act_vol_row}*{col}{fcst_aht_row})/({col}{work_hours_row}*3600*{col}{occ_row}*(1-{col}{fcst_shrink_row})),0)'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 1

    fcst_vA = row
    ws[f'A{row}'] = 'v(A) - AHT at Actual'
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'=IFERROR(({col}{fcst_vol_row}*{col}{act_aht_row})/({col}{work_hours_row}*3600*{col}{occ_row}*(1-{col}{fcst_shrink_row})),0)'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 1

    fcst_vS = row
    ws[f'A{row}'] = 'v(S) - Shrink at Actual'
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'=IFERROR(({col}{fcst_vol_row}*{col}{fcst_aht_row})/({col}{work_hours_row}*3600*{col}{occ_row}*(1-{col}{act_shrink_row})),0)'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 1

    fcst_vVA = row
    ws[f'A{row}'] = 'v(V,A)'
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'=IFERROR(({col}{act_vol_row}*{col}{act_aht_row})/({col}{work_hours_row}*3600*{col}{occ_row}*(1-{col}{fcst_shrink_row})),0)'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 1

    fcst_vVS = row
    ws[f'A{row}'] = 'v(V,S)'
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'=IFERROR(({col}{act_vol_row}*{col}{fcst_aht_row})/({col}{work_hours_row}*3600*{col}{occ_row}*(1-{col}{act_shrink_row})),0)'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 1

    fcst_vAS = row
    ws[f'A{row}'] = 'v(A,S)'
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'=IFERROR(({col}{fcst_vol_row}*{col}{act_aht_row})/({col}{work_hours_row}*3600*{col}{occ_row}*(1-{col}{act_shrink_row})),0)'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 1

    fcst_vVAS = row
    ws[f'A{row}'] = 'v(V,A,S) - All at Actual'
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'={col}{act_fte_row}'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 2

    # Shapley Attribution
    ws[f'A{row}'] = 'Shapley Attribution (FTE)'
    ws[f'A{row}'].font = Font(bold=True)
    row += 1

    fcst_phiV = row
    ws[f'A{row}'] = 'φ(Volume)'
    ws[f'A{row}'].font = Font(bold=True)
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'=(1/3)*({col}{fcst_vV}-{col}{fcst_v0})+(1/6)*({col}{fcst_vVA}-{col}{fcst_vA})+(1/6)*({col}{fcst_vVS}-{col}{fcst_vS})+(1/3)*({col}{fcst_vVAS}-{col}{fcst_vAS})'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 1

    fcst_phiA = row
    ws[f'A{row}'] = 'φ(AHT)'
    ws[f'A{row}'].font = Font(bold=True)
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'=(1/3)*({col}{fcst_vA}-{col}{fcst_v0})+(1/6)*({col}{fcst_vVA}-{col}{fcst_vV})+(1/6)*({col}{fcst_vAS}-{col}{fcst_vS})+(1/3)*({col}{fcst_vVAS}-{col}{fcst_vVS})'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 1

    fcst_phiS = row
    ws[f'A{row}'] = 'φ(Shrinkage)'
    ws[f'A{row}'].font = Font(bold=True)
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'=(1/3)*({col}{fcst_vS}-{col}{fcst_v0})+(1/6)*({col}{fcst_vVS}-{col}{fcst_vV})+(1/6)*({col}{fcst_vAS}-{col}{fcst_vA})+(1/3)*({col}{fcst_vVAS}-{col}{fcst_vVA})'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 1

    ws[f'A{row}'] = 'Check (should = variance)'
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'={col}{fcst_phiV}+{col}{fcst_phiA}+{col}{fcst_phiS}'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 3

    # ========== SHAPLEY: PLAN vs ACTUAL ==========
    ws[f'A{row}'] = 'SHAPLEY: PLAN vs ACTUAL'
    apply_section_style(ws[f'A{row}'])
    ws[f'A{row}'].fill = PatternFill(start_color=COLORS['gold'], end_color=COLORS['gold'], fill_type='solid')
    ws[f'A{row}'].font = Font(bold=True, size=12, color=COLORS['header'])
    row += 1

    ws[f'A{row}'] = 'Coalition values (baseline = plan/budget values)'
    ws[f'A{row}'].font = Font(italic=True, size=10)
    row += 1

    # Coalition values for Plan
    plan_v0 = row
    ws[f'A{row}'] = 'v(∅) - All at Plan'
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'={col}{plan_fte_row}'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 1

    plan_vV = row
    ws[f'A{row}'] = 'v(V) - Volume at Actual'
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'=IFERROR(({col}{act_vol_row}*{col}{plan_aht_row})/({col}{work_hours_row}*3600*{col}{occ_row}*(1-{col}{plan_shrink_row})),0)'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 1

    plan_vA = row
    ws[f'A{row}'] = 'v(A) - AHT at Actual'
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'=IFERROR(({col}{plan_vol_row}*{col}{act_aht_row})/({col}{work_hours_row}*3600*{col}{occ_row}*(1-{col}{plan_shrink_row})),0)'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 1

    plan_vS = row
    ws[f'A{row}'] = 'v(S) - Shrink at Actual'
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'=IFERROR(({col}{plan_vol_row}*{col}{plan_aht_row})/({col}{work_hours_row}*3600*{col}{occ_row}*(1-{col}{act_shrink_row})),0)'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 1

    plan_vVA = row
    ws[f'A{row}'] = 'v(V,A)'
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'=IFERROR(({col}{act_vol_row}*{col}{act_aht_row})/({col}{work_hours_row}*3600*{col}{occ_row}*(1-{col}{plan_shrink_row})),0)'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 1

    plan_vVS = row
    ws[f'A{row}'] = 'v(V,S)'
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'=IFERROR(({col}{act_vol_row}*{col}{plan_aht_row})/({col}{work_hours_row}*3600*{col}{occ_row}*(1-{col}{act_shrink_row})),0)'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 1

    plan_vAS = row
    ws[f'A{row}'] = 'v(A,S)'
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'=IFERROR(({col}{plan_vol_row}*{col}{act_aht_row})/({col}{work_hours_row}*3600*{col}{occ_row}*(1-{col}{act_shrink_row})),0)'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 1

    plan_vVAS = row
    ws[f'A{row}'] = 'v(V,A,S) - All at Actual'
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'={col}{act_fte_row}'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 2

    # Shapley Attribution for Plan
    ws[f'A{row}'] = 'Shapley Attribution (FTE)'
    ws[f'A{row}'].font = Font(bold=True)
    row += 1

    plan_phiV = row
    ws[f'A{row}'] = 'φ(Volume)'
    ws[f'A{row}'].font = Font(bold=True)
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'=(1/3)*({col}{plan_vV}-{col}{plan_v0})+(1/6)*({col}{plan_vVA}-{col}{plan_vA})+(1/6)*({col}{plan_vVS}-{col}{plan_vS})+(1/3)*({col}{plan_vVAS}-{col}{plan_vAS})'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 1

    plan_phiA = row
    ws[f'A{row}'] = 'φ(AHT)'
    ws[f'A{row}'].font = Font(bold=True)
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'=(1/3)*({col}{plan_vA}-{col}{plan_v0})+(1/6)*({col}{plan_vVA}-{col}{plan_vV})+(1/6)*({col}{plan_vAS}-{col}{plan_vS})+(1/3)*({col}{plan_vVAS}-{col}{plan_vVS})'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 1

    plan_phiS = row
    ws[f'A{row}'] = 'φ(Shrinkage)'
    ws[f'A{row}'].font = Font(bold=True)
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'=(1/3)*({col}{plan_vS}-{col}{plan_v0})+(1/6)*({col}{plan_vVS}-{col}{plan_vV})+(1/6)*({col}{plan_vAS}-{col}{plan_vA})+(1/3)*({col}{plan_vVAS}-{col}{plan_vVA})'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 1

    ws[f'A{row}'] = 'Check (should = variance)'
    for m in range(12):
        col = get_column_letter(m+2)
        ws.cell(row=row, column=m+2).value = f'={col}{plan_phiV}+{col}{plan_phiA}+{col}{plan_phiS}'
        ws.cell(row=row, column=m+2).number_format = '#,##0.00'
    row += 3

    # ========== METHODOLOGY ==========
    ws[f'A{row}'] = 'METHODOLOGY'
    apply_section_style(ws[f'A{row}'])
    row += 1

    notes = [
        "Shapley value decomposition treats variance attribution as a cooperative game.",
        "Each factor's contribution = average marginal impact across all possible orderings.",
        "For n=3 factors, we compute 2³=8 coalition values and apply weights: 1/3, 1/6, 1/6, 1/3.",
        "",
        "Key assumptions:",
        "• Occupancy is a PLANNING CONSTANT - it does not contribute to variance",
        "• Baseline for 'Fcst vs Actual' = Forecast values",
        "• Baseline for 'Plan vs Actual' = Plan/Budget values",
        "",
        "Formula: FTE = (Volume × AHT) / (WorkHours × 3600 × Occupancy × (1 - Shrinkage))",
        "• Volume: Number of contacts/calls",
        "• AHT: Average Handle Time in SECONDS",
        "• WorkHours: Productive hours per month",
        "• Shrinkage: Fraction of time agents are unavailable",
        "",
        "Shapley values sum EXACTLY to total variance (verify with 'Check' row).",
    ]

    for note in notes:
        ws[f'A{row}'] = note
        ws[f'A{row}'].font = Font(size=10)
        row += 1

    # Freeze panes
    ws.freeze_panes = 'B4'

    # Save
    wb.save(output_path)
    print(f"✅ WFM Shapley template generated: {output_path}")
    print(f"   Year: {year}")
    print(f"   Factors: Volume, AHT, Shrinkage (Occupancy held constant)")
    print(f"   Analyses: Forecast vs Actual, Plan vs Actual")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description='Generate WFM Shapley Decomposition Template')
    parser.add_argument('--output', '-o', required=True, help='Output Excel file path')
    parser.add_argument('--year', '-y', type=int, default=2025, help='Year for column headers')
    args = parser.parse_args()

    generate_wfm_shapley(args.output, args.year)
